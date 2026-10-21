# Analisador de Pagamentos — Python
import pandas as pd
import matplotlib.pyplot as plt


# ==========================================
# 1. CONFIGURAÇÃO E LEITURA DOS DADOS
# ==========================================

ARQUIVO = "Planilha de Pagamentos.xlsx"

print("=" * 45)
print("       ANALISADOR DE PAGAMENTOS")
print("=" * 45)

dados = pd.read_excel(ARQUIVO, header=2)


# ==========================================
# 2. LIMPEZA DOS DADOS
# ==========================================

# Remove colunas completamente vazias.
dados_limpos = dados.dropna(axis=1, how="all")

print("\n===== ESTRUTURA DOS DADOS =====")
print(f"Quantidade de registros: {len(dados_limpos)}")
print(f"Quantidade de colunas: {len(dados_limpos.columns)}")

print("\nColunas disponíveis:")
print(dados_limpos.columns.tolist())

print("\nValores ausentes por coluna:")
print(dados_limpos.isna().sum())

print("\nTipos de dados:")
print(dados_limpos.dtypes)

print("\nPrimeiros cinco registros:")
print(dados_limpos.head().to_string(index=False))


# ==========================================
# 3. FUNÇÃO DE FORMATAÇÃO MONETÁRIA
# ==========================================

def formatar_reais(valor):
    """Formata um número no padrão monetário brasileiro."""
    return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# ==========================================
# 4. RESUMO FINANCEIRO
# ==========================================

valores = dados_limpos["TOTAL BRUTO"]

print("\n===== RESUMO FINANCEIRO =====")
print(f"Total bruto:          {formatar_reais(valores.sum())}")
print(f"Média por pagamento:  {formatar_reais(valores.mean())}")
print(f"Maior pagamento:      {formatar_reais(valores.max())}")
print(f"Menor pagamento:      {formatar_reais(valores.min())}")
print(f"Total de registros:   {valores.count()}")


# ==========================================
# 5. DETALHES DO MAIOR PAGAMENTO
# ==========================================

indice_maior = valores.idxmax()
maior_pagamento = dados_limpos.loc[indice_maior]

print("\n===== MAIOR PAGAMENTO =====")
print(f"Cliente:              {maior_pagamento['CLIENTE']}")
print(f"Data:                 {maior_pagamento['DATA'].strftime('%d/%m/%Y')}")
print(f"Valor:                {formatar_reais(maior_pagamento['TOTAL BRUTO'])}")
print(f"Forma de pagamento:   {maior_pagamento['FORMA DE PAGAMENTO']}")


# ==========================================
# 6. ANÁLISE POR FORMA DE PAGAMENTO
# ==========================================

resumo_pagamentos = dados_limpos.groupby("FORMA DE PAGAMENTO").agg(
    quantidade=("TOTAL BRUTO", "count"),
    total=("TOTAL BRUTO", "sum"),
    media=("TOTAL BRUTO", "mean")
)

# Cria uma cópia para formatar os valores sem alterar
# os números do resumo original.
resumo_formatado = resumo_pagamentos.copy()

resumo_formatado["total"] = resumo_formatado["total"].apply(formatar_reais)
resumo_formatado["media"] = resumo_formatado["media"].apply(formatar_reais)

print("\n===== ANÁLISE POR FORMA DE PAGAMENTO =====")
print(resumo_formatado.to_string())

# ==========================================
# 7. GRÁFICO POR FORMA DE PAGAMENTO
# ==========================================

fig, ax = plt.subplots(figsize=(9, 5))

barras = ax.bar(
    resumo_pagamentos.index,
    resumo_pagamentos["total"]
)

ax.set_title("Total recebido por forma de pagamento")
ax.set_xlabel("Forma de pagamento")
ax.set_ylabel("Total recebido (R$)")

for barra in barras:
    valor = barra.get_height()

    ax.annotate(
        formatar_reais(valor),
        (barra.get_x() + barra.get_width() / 2, valor),
        ha="center",
        va="bottom",
        xytext=(0, 5),
        textcoords="offset points"
    )

ax.set_ylim(0, resumo_pagamentos["total"].max() * 1.15)

fig.tight_layout()
plt.show()

# ==========================================
# 8. RELATÓRIO EXCEL FORMATADO
# ==========================================

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ARQUIVO_RELATORIO = "Relatorio de Pagamentos.xlsx"

with pd.ExcelWriter(ARQUIVO_RELATORIO, engine="openpyxl") as escritor:

    # ABA 1: DADOS TRATADOS
    dados_limpos.to_excel(
        escritor,
        sheet_name="Dados tratados",
        index=False
    )

    # ABA 2: RESUMO POR FORMA DE PAGAMENTO
    resumo_pagamentos.to_excel(
        escritor,
        sheet_name="Por pagamento",
    
     index=True,
        index_label="Forma de pagamento"
    )

    # ABA 3: RELATÓRIO PRINCIPAL
    relatorio = escritor.book.create_sheet("Relatório", 0)

    relatorio.merge_cells("A1:B1")
    relatorio["A1"] = "RELATÓRIO FINANCEIRO DE PAGAMENTOS"

    relatorio["A3"] = "Indicador"
    relatorio["B3"] = "Valor"

    indicadores = [
        ("Total bruto recebido", valores.sum()),
        ("Média por pagamento", valores.mean()),
        ("Maior pagamento", valores.max()),
        ("Menor pagamento", valores.min()),
        ("Quantidade de pagamentos", valores.count()),
    ]

    for linha, (nome, valor) in enumerate(indicadores, start=4):
        relatorio.cell(row=linha, column=1, value=nome)
        relatorio.cell(row=linha, column=2, value=float(valor))

    # Detalhes do maior pagamento
    relatorio.merge_cells("A11:B11")
    relatorio["A11"] = "DETALHES DO MAIOR PAGAMENTO"

    detalhes = [
        ("Cliente", str(maior_pagamento["CLIENTE"])),
        ("Data", maior_pagamento["DATA"].to_pydatetime()),
        ("Forma de pagamento", str(maior_pagamento["FORMA DE PAGAMENTO"])),
        ("Valor", float(maior_pagamento["TOTAL BRUTO"])),
    ]

    for linha, (nome, valor) in enumerate(detalhes, start=12):
        relatorio.cell(row=linha, column=1, value=nome)
        relatorio.cell(row=linha, column=2, value=valor)

    # CORES E ESTILOS
    azul_escuro = "17365D"
    azul_claro = "D9EAF7"
    branco = "FFFFFF"
    cinza = "D9E1F2"

    preenchimento_titulo = PatternFill(
        fill_type="solid", fgColor=azul_escuro
    )
    preenchimento_cabecalho = PatternFill(
        fill_type="solid", fgColor=azul_claro
    )
    preenchimento_secao = PatternFill(
        fill_type="solid", fgColor=cinza
    )

    borda_fina = Border(
        bottom=Side(style="thin", color="D9E1F2")
    )

    # Formatação da aba Relatório
    for celula in relatorio[1]:
        celula.fill = preenchimento_titulo
        celula.font = Font(color=branco, bold=True, size=16)
        celula.alignment = Alignment(
            horizontal="center", vertical="center"
        )

    relatorio.row_dimensions[1].height = 32

    for celula in relatorio[3]:
        celula.fill = preenchimento_cabecalho
        celula.font = Font(bold=True, color=azul_escuro)
        celula.alignment = Alignment(horizontal="center")

    for celula in relatorio[11]:
        celula.fill = preenchimento_titulo
        celula.font = Font(color=branco, bold=True)

    for linha in range(4, 9):
        relatorio.cell(linha, 1).font = Font(bold=True)
        relatorio.cell(linha, 1).border = borda_fina
        relatorio.cell(linha, 2).border = borda_fina

    # Formatação monetária no Excel
    formato_reais = '[$R$-416] #,##0.00'

    for linha in range(4, 8):
        relatorio.cell(linha, 2).number_format = formato_reais

    relatorio["B8"].number_format = "0"

    for linha in range(12, 16):
        relatorio.cell(linha, 1).font = Font(bold=True)
        relatorio.cell(linha, 1).border = borda_fina
        relatorio.cell(linha, 2).border = borda_fina

    relatorio["B13"].number_format = "DD/MM/YYYY"
    relatorio["B15"].number_format = formato_reais

    # Formatação das outras duas abas
    for nome_aba in ["Dados tratados", "Por pagamento"]:
        planilha = escritor.sheets[nome_aba]

        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = planilha.dimensions

        for celula in planilha[1]:
            celula.fill = preenchimento_titulo
            celula.font = Font(color=branco, bold=True)
            celula.alignment = Alignment(
                horizontal="center", vertical="center"
            )

        for coluna in planilha.columns:
            letra = coluna[0].column_letter
            maior_largura = max(
                len(str(celula.value or ""))
                for celula in coluna
            )
            planilha.column_dimensions[letra].width = min(
                maior_largura + 3, 32
            )

    # Formatar valores monetários na aba de pagamentos
    pagamentos = escritor.sheets["Por pagamento"]

    for linha in range(2, pagamentos.max_row + 1):
        pagamentos.cell(linha, 2).number_format = "0"
        pagamentos.cell(linha, 3).number_format = formato_reais
        pagamentos.cell(linha, 4).number_format = formato_reais

    # Ajustes finais
    relatorio.column_dimensions["A"].width = 32
    relatorio.column_dimensions["B"].width = 24
    relatorio.freeze_panes = "A4"
    relatorio.sheet_view.showGridLines = False

print(f"\nRelatório formatado exportado: {ARQUIVO_RELATORIO}")
